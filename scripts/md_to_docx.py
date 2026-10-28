"""把论文 Markdown 转成符合中文学位论文习惯的 Word 文档。

处理范围（只覆盖论文实际用到的 Markdown 子集）：
    标题四级、段落、粗体/斜体/行内代码/行内公式、行间公式（LaTeX→OMML）、
    管道表格（三线表）、代码块（灰底框）、有序/无序列表、引用块、图片

版式约定：
    A4，上 2.5 / 下 2.5 / 左 3.0 / 右 2.5 cm，正文宽 15.5 cm
    正文   宋体小四（12pt）／西文 Times New Roman／1.5 倍行距／首行缩进 2 字符
    章标题 黑体三号（16pt）居中，段前分页
    节标题 黑体四号（14pt）左对齐
    小节   黑体小四（12pt）左对齐
    图题表题 宋体五号（10.5pt）居中
    表格   宋体五号，三线表
    代码   Consolas 小五（9pt），浅灰底

用法：
    python scripts/md_to_docx.py [输入.md] [输出.docx]
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import omml  # noqa: E402

# ---------------------------------------------------------------- 版式常量

SONG, HEI, KAI, MONO, LATIN = "宋体", "黑体", "楷体", "Consolas", "Times New Roman"
BODY_PT = 12.0          # 小四
CAPTION_PT = 10.5       # 五号
TABLE_PT = 10.5
CODE_PT = 9.0           # 小五
TEXT_W_CM = 15.5
BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x44, 0x44, 0x44)

CM2TWIP = 1440 / 2.54

# OOXML 对属性的子元素顺序有硬性要求，顺序错了 Word 会拒绝打开，
# 因此新增子元素一律按 schema 声明的序列插入，不能 append 了事。
PPR_ORDER = [
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
    "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
    "w:sectPr", "w:pPrChange",
]
TCPR_ORDER = [
    "w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders",
    "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
    "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange",
]
TBLPR_ORDER = [
    "w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
    "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd",
    "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
    "w:tblCaption", "w:tblDescription", "w:tblPrChange",
]


def tw(cm: float) -> int:
    return int(round(cm * CM2TWIP))


def _insert_ordered(parent, el, order: list[str]) -> None:
    """按 schema 顺序把 el 插入 parent。"""
    if el.tag not in [qn(t) for t in order]:
        parent.append(el)
        return
    # insert_element_before 内部会对标签名再做一次 qn()，因此这里传原始前缀名
    successors = order[order.index(next(t for t in order if qn(t) == el.tag)) + 1:]
    parent.insert_element_before(el, *successors)


def _replace_ordered(parent, el, order: list[str]) -> None:
    existing = parent.find(el.tag)
    if existing is not None:
        parent.remove(existing)
    _insert_ordered(parent, el, order)


# ---------------------------------------------------------------- 字体工具


def _set_rfonts(rpr, latin: str, ea: str) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:cs"), latin)
    rfonts.set(qn("w:eastAsia"), ea)


def run_font(run, latin: str, ea: str) -> None:
    _set_rfonts(run._element.get_or_add_rPr(), latin, ea)


def style_font(style, latin: str, ea: str, size: float | None = None,
               bold: bool | None = None, color: RGBColor | None = None) -> None:
    _set_rfonts(style.element.get_or_add_rPr(), latin, ea)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def no_indent(paragraph) -> None:
    """清零缩进（同时清掉字符单位的缩进，避免覆盖磅值）。"""
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        _insert_ordered(pPr, ind, PPR_ORDER)
    for attr in ("w:firstLineChars", "w:leftChars", "w:rightChars"):
        ind.attrib.pop(qn(attr), None)
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:left"), "0")


def first_line_chars(paragraph, chars: int = 200, pt: float = BODY_PT) -> None:
    """首行缩进 2 字符：同时写 firstLineChars（中文习惯）与 firstLine（兼容）。"""
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        _insert_ordered(pPr, ind, PPR_ORDER)
    ind.set(qn("w:firstLineChars"), str(chars))
    ind.set(qn("w:firstLine"), str(int(round(chars / 100 * pt * 20))))


def set_ppr(pPr, el) -> None:
    _insert_ordered(pPr, el, PPR_ORDER)


# ---------------------------------------------------------------- 表格修饰


def shade_cell(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    _replace_ordered(tcPr, shd, TCPR_ORDER)


def cell_borders(cell, **edges: tuple[str, int]) -> None:
    """edges 例：top=("single", 12)；sz 单位为 1/8 磅。"""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        _insert_ordered(tcPr, borders, TCPR_ORDER)
    for edge in ("top", "left", "start", "bottom", "right", "end", "insideH", "insideV"):
        if edge not in edges:
            continue
        val, sz = edges[edge]
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)


def table_three_line(table, header_rows: int = 1) -> None:
    """三线表：顶线粗、表头下细线、底线粗，无竖线、无内横线。"""
    tblPr = table._tbl.tblPr

    # 单元格默认左右边距各 0.19 cm，八列以上时会吃掉近 0.4 cm 可用宽度，
    # 导致「0.294±0.027」这类数字在中间被迫换行。三线表无需留白，收到 0.08 cm。
    cell_mar = OxmlElement("w:tblCellMar")
    for side, width in (("top", 20), ("left", 45), ("bottom", 20), ("right", 45)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(width))
        el.set(qn("w:type"), "dxa")
        cell_mar.append(el)
    _replace_ordered(tblPr, cell_mar, TBLPR_ORDER)

    borders = OxmlElement("w:tblBorders")
    for edge, val, sz in (
        ("top", "single", 12),
        ("left", "none", 0),
        ("bottom", "single", 12),
        ("right", "none", 0),
        ("insideH", "none", 0),
        ("insideV", "none", 0),
    ):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    _replace_ordered(tblPr, borders, TBLPR_ORDER)

    for cell in table.rows[header_rows - 1].cells:
        cell_borders(cell, bottom=("single", 6))


def set_col_widths(table, weights: list[float], total_cm: float = TEXT_W_CM) -> None:
    """设定列宽。

    固定布局（tblLayout=fixed）下 Word 按 tblGrid 的 gridCol 排版，
    只改单元格的 tcW 不生效，因此列宽与单元格宽度必须同时写。
    """
    widths = [max(1.4, total_cm * w / (sum(weights) or 1)) for w in weights]
    scale = total_cm / sum(widths)
    widths = [w * scale for w in widths]
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)


def add_field(paragraph, instr: str, placeholder: str = "") -> None:
    """插入 Word 域（PAGE / TOC）。"""
    for kind in ("begin", "instr", "separate", "text", "end"):
        run = paragraph.add_run()
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
            run._r.append(el)
        elif kind == "text":
            if placeholder:
                run.text = placeholder
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)


# ---------------------------------------------------------------- 行内解析

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\$[^$]+\$|\*[^*\s][^*]*\*)")


def add_inline(paragraph, text: str, *, latin: str = LATIN, ea: str = SONG,
               size: float = BODY_PT, bold: bool = False, italic: bool = False,
               color: RGBColor | None = None, keep_font: bool = False,
               stats: dict | None = None) -> None:
    def emit(chunk: str, *, is_code=False, is_bold=False, is_italic=False):
        if not chunk:
            return
        run = paragraph.add_run(chunk)
        if not keep_font:
            run_font(run, MONO if is_code else latin, ea)
            run.font.size = Pt(size - 0.5 if is_code else size)
        if is_bold or bold:
            run.bold = True
        if is_italic or italic:
            run.italic = True
        if color is not None:
            run.font.color.rgb = color

    def emit_math(src: str) -> None:
        try:
            paragraph._p.append(parse_xml(omml.omml_element(src)))
            if stats is not None:
                stats["math"] = stats.get("math", 0) + 1
        except Exception:
            plain_src = re.sub(r"\\[a-zA-Z]+|[{}$\\]", "", src).strip()
            emit(plain_src or src)
            if stats is not None:
                stats["math_fallback"] = stats.get("math_fallback", 0) + 1

    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            emit(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            add_inline(paragraph, tok[2:-2], latin=latin, ea=ea, size=size, bold=True,
                       italic=italic, color=color, keep_font=keep_font, stats=stats)
        elif tok.startswith("`"):
            emit(tok[1:-1], is_code=True)
        elif tok.startswith("$"):
            emit_math(tok[1:-1])
        else:
            add_inline(paragraph, tok[1:-1], latin=latin, ea=ea, size=size, bold=bold,
                       italic=True, color=color, keep_font=keep_font, stats=stats)
        pos = m.end()
    if pos < len(text):
        emit(text[pos:])


def plain(text: str) -> str:
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    s = re.sub(r"`(.+?)`", r"\1", s)
    s = re.sub(r"\$(.+?)\$", r"\1", s)
    s = re.sub(r"\*(.+?)\*", r"\1", s)
    return re.sub(r"\\[a-zA-Z]+", "", s)


def disp_width(text: str) -> int:
    return sum(2 if ord(c) > 0x2E80 else 1 for c in plain(text))


# ---------------------------------------------------------------- 文档骨架


def build_document() -> Document:
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.5)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.5)
    sec.header_distance, sec.footer_distance = Cm(1.5), Cm(1.75)

    normal = doc.styles["Normal"]
    style_font(normal, LATIN, SONG, BODY_PT, bold=False, color=BLACK)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, center, before, after in (
        ("Heading 1", 16, True, 6, 18),
        ("Heading 2", 14, False, 12, 6),
        ("Heading 3", 12, False, 6, 3),
    ):
        st = doc.styles[name]
        style_font(st, LATIN, HEI, size, bold=True, color=BLACK)
        pf = st.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.line_spacing = 1.5
        pf.first_line_indent = Pt(0)
        pf.keep_with_next = True
        if name == "Heading 1":
            pf.page_break_before = True

    title = doc.styles.add_style("论文题目", WD_STYLE_TYPE.PARAGRAPH)
    title.base_style = doc.styles["Normal"]
    style_font(title, LATIN, HEI, 22, bold=True, color=BLACK)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(105)
    title.paragraph_format.space_after = Pt(14)
    title.paragraph_format.first_line_indent = Pt(0)

    for style_name, size, ea, bold in (("图表题", CAPTION_PT, SONG, False),
                                       ("目录标题", 16, HEI, True)):
        st = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        style_font(st, LATIN, ea, size, bold=bold, color=BLACK)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st.paragraph_format.first_line_indent = Pt(0)
        st.paragraph_format.line_spacing = 1.25
        st.paragraph_format.space_before = Pt(6)
        st.paragraph_format.space_after = Pt(18 if style_name == "目录标题" else 12)
        if style_name == "目录标题":
            st.paragraph_format.page_break_before = True

    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer_p, "PAGE", "1")
    for run in footer_p.runs:
        run_font(run, LATIN, SONG)
        run.font.size = Pt(CAPTION_PT)

    return doc


# ---------------------------------------------------------------- 块级渲染


class Writer:
    def __init__(self, doc: Document, md_dir: Path) -> None:
        self.doc = doc
        self.md_dir = md_dir
        self.stats = {"para": 0, "table": 0, "image": 0, "code": 0, "math": 0,
                      "math_fallback": 0, "heading": 0}
        self.seen_title = False
        self.toc_done = False

    # -- 段落
    def para(self, text: str, *, indent: bool = True, align=None, ea: str = SONG,
             size: float = BODY_PT, color: RGBColor | None = None):
        p = self.doc.add_paragraph()
        if indent:
            first_line_chars(p, 200, size)
        else:
            no_indent(p)
        p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        # 表题在表格之前，必须与表格同页；图题在图片之后，图片已设 keep_with_next。
        if re.match(r"^\*\*表\s*\d+[-–]\d+", text.strip()):
            p.paragraph_format.keep_with_next = True
        add_inline(p, text, ea=ea, size=size, color=color, stats=self.stats)
        self.stats["para"] += 1
        return p

    def subtitle(self, text: str) -> None:
        p = self.doc.add_paragraph()
        no_indent(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(30)
        add_inline(p, text, ea=KAI, size=14, color=GREY, stats=self.stats)

    # -- 标题
    def heading(self, level: int, text: str) -> None:
        if not self.seen_title and level == 1:
            p = self.doc.add_paragraph(style="论文题目")
            add_inline(p, text, keep_font=True, stats=self.stats)
            self.seen_title = True
            self.stats["heading"] += 1
            return

        if not self.toc_done and re.match(r"^第\s*1\s*章", text):
            self.add_toc()

        p = self.doc.add_paragraph(style=f"Heading {min(level, 3)}")
        add_inline(p, text, keep_font=True, stats=self.stats)
        self.stats["heading"] += 1

    def add_toc(self) -> None:
        head = self.doc.add_paragraph(style="目录标题")
        run = head.add_run("目　　录")
        run_font(run, LATIN, HEI)
        run.font.size = Pt(16)

        p = self.doc.add_paragraph()
        no_indent(p)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.5
        add_field(p, 'TOC \\o "1-3" \\h \\z \\u',
                  "【目录：在 Word 中右键 → 更新域，即可自动生成】")
        self.toc_done = True

    def caption(self, text: str, *, keep_next: bool = False) -> None:
        p = self.doc.add_paragraph(style="图表题")
        add_inline(p, text, ea=SONG, size=CAPTION_PT, stats=self.stats)
        p.paragraph_format.keep_with_next = keep_next

    # -- 行间公式
    def display_math(self, latex: str) -> None:
        m = re.search(r"\\tag\{([^}]*)\}", latex)
        tag = m.group(1).strip() if m else None
        src = re.sub(r"\\tag\{[^}]*\}", "", latex).strip()

        p = self.doc.add_paragraph()
        no_indent(p)
        pf = p.paragraph_format
        pf.line_spacing = 1.5
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)

        if tag is None:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            pPr = p._p.get_or_add_pPr()
            tabs = OxmlElement("w:tabs")
            for val, pos in (("center", tw(TEXT_W_CM / 2)), ("right", tw(TEXT_W_CM))):
                tab = OxmlElement("w:tab")
                tab.set(qn("w:val"), val)
                tab.set(qn("w:pos"), str(pos))
                tabs.append(tab)
            set_ppr(pPr, tabs)
            p.add_run("\t")

        try:
            p._p.append(parse_xml(omml.omml_element(src)))
            self.stats["math"] += 1
        except Exception as exc:
            run = p.add_run(f"[公式解析失败：{exc}] {src}")
            run_font(run, MONO, SONG)
            run.font.size = Pt(CODE_PT)
            self.stats["math_fallback"] += 1

        if tag:
            run = p.add_run(f"\t（{tag}）")
            run_font(run, LATIN, SONG)
            run.font.size = Pt(BODY_PT)

    # -- 代码块 / 伪代码
    def code_block(self, lines: list[str]) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Cm(TEXT_W_CM)
        shade_cell(cell, "F7F7F9")
        cell_borders(cell, top=("single", 4), left=("single", 4),
                     bottom=("single", 4), right=("single", 4))
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        first = True
        for line in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            no_indent(p)
            pf = p.paragraph_format
            pf.line_spacing = 1.0
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            run = p.add_run(line if line.strip() else " ")
            run_font(run, MONO, SONG)
            run.font.size = Pt(CODE_PT)
        self.stats["code"] += 1
        self.spacer(4)

    # -- 图片
    def image(self, path: Path, alt: str) -> None:
        if not path.exists():
            self.para(f"［缺失图片：{path.name}］", color=RGBColor(0xC0, 0, 0))
            return
        p = self.doc.add_paragraph()
        no_indent(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(TEXT_W_CM - 0.1))
        self.stats["image"] += 1

    # -- 表格
    def table(self, raw_rows: list[str]) -> None:
        rows = [r for r in raw_rows if r.strip()]
        rows = [r for r in rows if not re.match(r"^\|[\s:\-|]+\|?\s*$", r.strip())]
        if not rows:
            return
        grid = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        ncol = max(len(r) for r in grid)
        grid = [r + [""] * (ncol - len(r)) for r in grid]

        table = self.doc.add_table(rows=len(grid), cols=ncol)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        tblW = OxmlElement("w:tblW")
        tblW.set(qn("w:w"), str(tw(TEXT_W_CM)))
        tblW.set(qn("w:type"), "dxa")
        _replace_ordered(table._tbl.tblPr, tblW, TBLPR_ORDER)

        # 列宽权重：按各列最长内容估算，并用幂次压低极端比例，
        # 避免某一列独占版面；上限 60 防止长文本列把其他列挤扁。
        col_max = [max(disp_width(r[j]) for r in grid) for j in range(ncol)]
        weights = [min(max(m ** 0.85, 5), 60) for m in col_max]

        # 字号随内容自动降档：半角单位宽度约 S/2 pt，故一列所需宽度 ≈
        # disp × S × 0.01764 cm。取满足总宽度的最大档位，列数多但内容短的表
        # （如 10 列的基线对比）不该被无故缩小，而含「均值±标准差」的宽表必须缩。
        usable = TEXT_W_CM - ncol * 0.16          # 减去单元格左右边距
        need_per_pt = 0.01764 * sum(col_max)      # 每 1pt 字号所需总宽
        best = TABLE_PT - 3.0
        for size in (10.5, 10.0, 9.5, 9.0, 8.5, 8.0, 7.5):
            if need_per_pt * size <= usable * 0.97:
                best = size
                break
        table_pt = best

        # 对齐按列统一：长文本列左对齐，其余居中。逐格判断会导致同列参差不齐。
        col_left = [max(disp_width(r[j]) for r in grid) > 40 for j in range(ncol)]

        for i, row in enumerate(grid):
            for j, text in enumerate(row):
                cell = table.cell(i, j)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                no_indent(p)
                pf = p.paragraph_format
                pf.line_spacing = 1.0
                pf.space_before = Pt(1.5)
                pf.space_after = Pt(1.5)
                pf.alignment = (WD_ALIGN_PARAGRAPH.LEFT
                                if (i > 0 and col_left[j]) else WD_ALIGN_PARAGRAPH.CENTER)
                add_inline(p, text, ea=SONG, size=table_pt, bold=(i == 0), stats=self.stats)

        table_three_line(table, header_rows=1)
        set_col_widths(table, weights, TEXT_W_CM)

        # 表格不允许跨页断开：逐行给单元格段落设 keep_with_next，最后一行除外。
        # 否则 Word 会把表拆成「表题 + 首行」与「续表」，中间隔一个分页符，
        # 读起来像两张表。仅对常规行数的表生效，超长表仍允许自然分页。
        if len(grid) <= 16:
            for row in table.rows[:-1]:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True

        trPr = table.rows[0]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))

        self.stats["table"] += 1
        self.spacer(6)

    # -- 列表
    def list_item(self, text: str, marker: str) -> None:
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 1.5
        pf.left_indent = Cm(0.74)
        pf.first_line_indent = Cm(-0.74)
        pPr = p._p.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "left")
        tab.set(qn("w:pos"), str(tw(0.74)))
        tabs.append(tab)
        set_ppr(pPr, tabs)
        add_inline(p, f"{marker}\t{text}", ea=SONG, size=BODY_PT, stats=self.stats)
        self.stats["para"] += 1

    # -- 引用块
    def quote(self, lines: list[str]) -> None:
        for line in lines:
            p = self.doc.add_paragraph()
            no_indent(p)
            pf = p.paragraph_format
            pf.line_spacing = 1.4
            pf.left_indent = Cm(0.8)
            pf.right_indent = Cm(0.5)
            pf.space_before = Pt(3)
            pf.space_after = Pt(3)
            add_inline(p, line, ea=KAI, size=CAPTION_PT, color=GREY, stats=self.stats)
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "12")
            left.set(qn("w:space"), "8")
            left.set(qn("w:color"), "BBBBBB")
            pBdr.append(left)
            set_ppr(pPr, pBdr)
            self.stats["para"] += 1

    def spacer(self, pt: float = 6) -> None:
        p = self.doc.add_paragraph()
        no_indent(p)
        pf = p.paragraph_format
        pf.line_spacing = 1.0
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        run = p.add_run(" ")
        run.font.size = Pt(pt)


# ---------------------------------------------------------------- 主流程

BLOCK_START = re.compile(r"^(#{1,6}\s|```|\$\$|\||>|[-*]\s+|\d+[.)]\s+|!\[|-{3,}\s*$)")


def convert(md_path: Path, out_path: Path) -> dict:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = build_document()
    w = Writer(doc, md_path.parent)

    i, n = 0, len(lines)
    while i < n:
        line = lines[i].rstrip()
        s = line.strip()

        if not s:
            i += 1
            continue

        # 代码块
        if s.startswith("```"):
            i += 1
            buf: list[str] = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i].rstrip())
                i += 1
            i += 1
            w.code_block(buf)
            continue

        # 行间公式
        if s.startswith("$$"):
            if s == "$$":
                i += 1
                buf = []
                while i < n and lines[i].strip() != "$$":
                    buf.append(lines[i].strip())
                    i += 1
                i += 1
                w.display_math(" ".join(buf))
            else:
                body = s[2:-2] if s.endswith("$$") else s[2:]
                w.display_math(body)
                i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level >= 2 and not re.match(r"^\d+(\.\d+)*\s", text):
                level -= 1          # 摘要 / 参考文献 / 附录 视为章级
            w.heading(level, text)
            i += 1
            continue

        # 表格
        if s.startswith("|"):
            buf = []
            while i < n and lines[i].strip().startswith("|"):
                buf.append(lines[i])
                i += 1
            w.table(buf)
            continue

        # 图片
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", s)
        if m:
            w.image(md_path.parent / m.group(2), m.group(1))
            i += 1
            continue

        # 引用
        if s.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^>\s?", "", lines[i].strip()))
                i += 1
            w.quote(buf)
            continue

        # 分隔线
        if re.match(r"^-{3,}\s*$", s):
            i += 1
            continue

        # 列表
        if re.match(r"^[-*]\s+", s):
            while i < n and re.match(r"^[-*]\s+", lines[i].strip()):
                w.list_item(re.sub(r"^[-*]\s+", "", lines[i].strip()), "·")
                i += 1
            continue
        if re.match(r"^\d+[.)]\s+", s):
            while i < n and re.match(r"^\d+[.)]\s+", lines[i].strip()):
                mm = re.match(r"^(\d+)[.)]\s+(.*)$", lines[i].strip())
                w.list_item(mm.group(2), f"{mm.group(1)}.")
                i += 1
            continue

        # 图题 / 表题
        m = re.match(r"^\*\*((?:图|表)\s*\d+[.\-]\d+[^*]*)\*\*$", s)
        if m:
            w.caption(m.group(1).strip(), keep_next=m.group(1).startswith("表"))
            i += 1
            continue

        # 版本行
        if s.startswith("**") and "论文初稿" in s:
            w.subtitle(s)
            i += 1
            continue

        # 普通段落
        buf = [s]
        i += 1
        while i < n and lines[i].strip() and not BLOCK_START.match(lines[i].strip()):
            buf.append(lines[i].strip())
            i += 1
        w.para(" ".join(buf))

    doc.core_properties.title = "基于代价敏感 CatBoost 的 MOOC 辍学风险预测研究"
    doc.core_properties.subject = "本科毕业设计（论文）初稿"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)

    w.stats["paragraph_total"] = len(doc.paragraphs)
    w.stats["tables_total"] = len(doc.tables)
    return w.stats


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else next((root / "docs").glob("论文初稿-*.md"))
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")

    stats = convert(src, dst)
    print(f"输出：{dst}")
    print(f"体积：{dst.stat().st_size / 1024:.1f} KB")
    for k, v in stats.items():
        print(f"  {k:16s} {v}")


if __name__ == "__main__":
    main()
